import os
from pathlib import Path
from docx import Document

# selenium is required for browser automation; provide a helpful message if missing
try:
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
except ImportError:
    raise ImportError("selenium is not installed. Run 'pip install selenium' and ensure the Edge WebDriver is available on PATH.")

import time
from utils import load_config


def chunk_text(text, size):
    words = text.split()
    chunks = []
    current = ''
    for w in words:
        if len(current) + len(w) + 1 > size:
            chunks.append(current.strip())
            current = ''
        current += w + ' '
    if current.strip():
        chunks.append(current.strip())
    return chunks


def open_copilot(config):
    options = webdriver.EdgeOptions()
    # keep browser open after script
    options.add_experimental_option("detach", True)
    driver = webdriver.Edge(options=options)
    driver.get(config['copilot_url'])
    input("Please log in to Copilot in the opened Edge window, then press Enter here to continue...")
    return driver


def query_copilot(driver, prompt):
    # NOTE: the CSS selectors below are placeholders; inspect the Copilot page to adjust them.
    # the input box is usually a textarea with aria-label="Ask a question" or similar.
    textbox = WebDriverWait(driver, 60).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, "textarea[aria-label*='question']"))
    )
    try:
        textbox.click()
    except Exception:
        pass
    # attempt to set the value via javascript in case send_keys is dropped
    driver.execute_script("arguments[0].value = arguments[1];", textbox, prompt)
    # give the page a moment to register the change
    time.sleep(0.5)
    textbox.send_keys(Keys.ENTER)
    # wait for an answer block to appear (selector may need adjustment)
    try:
        response_div = WebDriverWait(driver, 120).until(
            EC.visibility_of_element_located((By.CSS_SELECTOR, "div.copilot-response"))
        )
    except Exception as e:
        print(f"Response element not found: {e}")
        # capture current page html for debugging
        print(driver.page_source[:1000])
        raise
    return response_div.text


def process_file(file_path, config, driver):
    doc = Document(file_path)
    full_text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())

    chunks = chunk_text(full_text, config['chunk_size'])
    corrected_chunks = []

    for chunk in chunks:
        print("Sending chunk to Copilot (length", len(chunk), ")")
        corrected = query_copilot(driver, chunk)
        corrected_chunks.append(corrected)
        time.sleep(2)  # brief pause between requests

    final_text = '\n'.join(corrected_chunks)

    # Save as .md
    output_path = Path(config['output_dir']) / f"{file_path.stem}_copilot.md"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("# Copilot-corrected document\n\n")
        f.write(final_text)
    print(f"Saved Copilot result to {output_path}")


def main():
    config = load_config()
    workspace = Path(__file__).parent
    input_dir = workspace / config['input_dir']
    config['output_dir'] = str(workspace / config['output_dir'])

    if not input_dir.exists():
        print("Input directory not found; create it and add .docx files.")
        return

    driver = open_copilot(config)

    for file_path in input_dir.glob("*.docx"):
        process_file(file_path, config, driver)

if __name__ == '__main__':
    main()